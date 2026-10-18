"""Local catalog preparation and single, explicitly selected publisher asset downloads.

This is not a crawler: download URLs are supplied from observed publisher links.
No retry, account flow, discovery, or arbitrary site traversal is performed.
"""
import argparse
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
from datetime import datetime, timezone
from pathlib import Path

BASE = Path(__file__).resolve().parent
PRIVATE = BASE / 'private'
MANIFEST = PRIVATE / 'manifest.json'


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def read_manifest():
    return json.loads(MANIFEST.read_text()) if MANIFEST.exists() else []


def save_row(row):
    PRIVATE.mkdir(parents=True, exist_ok=True)
    rows = read_manifest()
    rows = [r for r in rows if r['id'] != row['id']]
    rows.append(row)
    MANIFEST.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n')


def download(args):
    old = next((r for r in read_manifest() if r['id'] == args.id), None)
    if old:
        print(json.dumps({'id': args.id, 'status': 'already_recorded', 'previous': old['status']}))
        return
    target = PRIVATE / 'sources' / args.family / args.filename
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        raise RuntimeError(f'Refusing overwrite: {target}')
    partial = target.with_suffix(target.suffix + '.part')
    cmd = ['curl', '--fail', '--location', '--connect-timeout', '20', '--max-time', '300',
           '--silent', '--show-error', '--output', str(partial), '--write-out', '%{json}', args.url]
    result = subprocess.run(cmd, capture_output=True, text=True)
    try:
        receipt = json.loads(result.stdout)
    except ValueError:
        receipt = {'raw': result.stdout[:500]}
    row = {'id': args.id, 'family': args.family, 'title': args.title, 'kind': args.kind,
           'source_page': args.source, 'asset_url': args.url, 'route': 'publisher_linked_asset_curl',
           'acquired_at': datetime.now(timezone.utc).isoformat(), 'receipt': receipt,
           'reading': 'not_read', 'status': 'failed', 'error': result.stderr[:1500]}
    valid = result.returncode == 0 and partial.exists() and partial.stat().st_size > 0
    if valid and target.suffix.lower() == '.pdf':
        with partial.open('rb') as f:
            valid = f.read(5) == b'%PDF-'
        if not valid:
            row['error'] = 'Response is not a PDF; retained as .part for diagnosis, not catalogued as a book.'
    if valid and target.suffix.lower() in ['.zip', '.docx']:
        with partial.open('rb') as f:
            valid = f.read(2) == b'PK'
        if not valid:
            row['error'] = 'Response is not a ZIP container.'
    if valid:
        partial.rename(target)
        row.update(status='downloaded', path=str(target), bytes=target.stat().st_size, sha256=sha(target))
    save_row(row)
    print(json.dumps({k: row.get(k) for k in ['id', 'status', 'bytes', 'error']}))


def inventory():
    keys = ['numenera', 'the strange', 'cypher system', 'arcana of the ancients', 'beasts of flesh',
            'in translation', 'the weird', 'planebreaker', 'ars magica', 'gurps', 'warhammer',
            'rifts', 'soulbound', 'maledictum', 'rogue trader', 'ironsworn', 'godbound', 'without number',
            'mythras', 'eclipse phase', 'pendragon', 'runequest', 'blades in the dark', 'harnmanor',
            'forbidden lands', 'power profiles', 'spheres of power', 'exalted', 'traveller',
            'swords of the serpentine', 'how to invent everything']
    roots = [Path('/Users/weiwei/Downloads'), Path('/Users/weiwei/Calibre Library')]
    existing = {r.get('original_path') for r in read_manifest()}
    count = 0
    for root in roots:
        for p in root.rglob('*'):
            if not p.is_file() or p.suffix.lower() not in ['.pdf', '.epub', '.mobi']:
                continue
            if not any(k in p.name.lower() for k in keys) or str(p) in existing:
                continue
            title = p.stem.split(' (')[0]
            file_id = 'local-' + hashlib.sha256(str(p).encode()).hexdigest()[:12]
            link = PRIVATE / 'existing' / (re.sub(r'[^a-zA-Z0-9]+', '-', title).strip('-') + p.suffix.lower())
            link.parent.mkdir(parents=True, exist_ok=True)
            if not link.exists():
                link.symlink_to(p)
            save_row({'id': file_id, 'family': 'existing-library', 'title': title, 'kind': 'existing_full_book',
                      'status': 'existing_linked', 'path': str(link), 'original_path': str(p),
                      'bytes': p.stat().st_size, 'sha256': sha(p), 'reading': 'not_read',
                      'source_page': None, 'route': 'local_inventory_symlink_original_untouched'})
            count += 1
    print(json.dumps({'new_existing_records': count}))


def prepare():
    from pypdf import PdfReader
    rows = read_manifest()
    db = PRIVATE / 'search.sqlite'
    con = sqlite3.connect(db)
    con.execute('CREATE VIRTUAL TABLE IF NOT EXISTS texts USING fts5(document_id UNINDEXED, title, locator UNINDEXED, body)')
    con.execute('DELETE FROM texts')
    pdftotext = shutil.which('pdftotext')
    if not pdftotext:
        candidate = Path('/Users/weiwei/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/fallback/pdftotext')
        pdftotext = str(candidate) if candidate.exists() else None
    text_root = PRIVATE / 'text'
    text_root.mkdir(exist_ok=True)
    for row in rows:
        if row['status'] not in ['downloaded', 'existing_linked', 'browser_downloaded']:
            continue
        p = Path(row['path'])
        if p.suffix.lower() == '.pdf':
            try:
                reader = PdfReader(p)
                row['pages'] = len(reader.pages)
                target = text_root / (row['id'] + '.txt')
                if not target.exists():
                    if pdftotext:
                        result = subprocess.run([pdftotext, '-layout', str(p), str(target)], capture_output=True, text=True)
                        if result.returncode:
                            raise RuntimeError(result.stderr[:300])
                    else:
                        target.write_text('\f'.join(page.extract_text() or '' for page in reader.pages))
                body = target.read_text(errors='replace')
                row.update(text_path=str(target), extracted_characters=len(body), preparation='parsed_and_text_extracted')
                for n, page in enumerate(body.split('\f'), 1):
                    if page.strip():
                        con.execute('INSERT INTO texts VALUES(?,?,?,?)', (row['id'], row['title'], 'page ' + str(n), page))
            except Exception as exc:
                row.update(preparation='parse_failed', preparation_error=str(exc)[:500])
        elif p.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(p)
            body = '\n'.join(v.text for v in doc.paragraphs)
            body += '\n' + '\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
            target = text_root / (row['id'] + '.txt')
            target.write_text(body)
            row.update(text_path=str(target), extracted_characters=len(body), preparation='parsed_and_text_extracted')
            con.execute('INSERT INTO texts VALUES(?,?,?,?)', (row['id'], row['title'], 'document', body))
    repo = PRIVATE / 'sources' / 'ars-magica-open-license'
    if (repo / '.git').exists():
        commit = subprocess.check_output(['git', '-C', str(repo), 'rev-parse', 'HEAD'], text=True).strip()
        ars = []
        for group in ['reviewed', 'wip', '3rd-party']:
            for p in sorted((repo / group).rglob('*.md')):
                body = p.read_text(errors='replace')
                book = {'path': str(p), 'title': p.stem, 'group': group, 'bytes': p.stat().st_size,
                        'sha256': sha(p), 'characters': len(body), 'reading': 'not_read'}
                ars.append(book)
                # Index sections, retaining exact source headings as private locators.
                for section in re.split(r'(?m)(?=^#{1,4} )', body):
                    if section.strip():
                        locator = section.split('\n', 1)[0][:250]
                        con.execute('INSERT INTO texts VALUES(?,?,?,?)', ('ars:' + str(p.relative_to(repo)), p.stem, locator, section))
        (PRIVATE / 'ars-magica-inventory.json').write_text(json.dumps({'commit': commit, 'books': ars}, ensure_ascii=False, indent=2) + '\n')
    con.commit()
    indexed = con.execute('SELECT COUNT(*) FROM texts').fetchone()[0]
    con.close()
    MANIFEST.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n')
    summary = {'records': len(rows), 'pdf_pages': sum(r.get('pages', 0) for r in rows), 'indexed_sections_or_pages': indexed,
               'failed_assets': [r['id'] for r in rows if r['status'] == 'failed'],
               'parse_failures': [r['id'] for r in rows if r.get('preparation') == 'parse_failed']}
    (PRIVATE / 'verification.json').write_text(json.dumps(summary, indent=2) + '\n')
    print(json.dumps(summary))


parser = argparse.ArgumentParser()
sub = parser.add_subparsers(dest='action', required=True)
sub.add_parser('inventory')
sub.add_parser('prepare')
dl = sub.add_parser('download')
for field in ['id', 'family', 'filename', 'title', 'kind', 'source', 'url']:
    dl.add_argument('--' + field, required=True)
args = parser.parse_args()
if args.action == 'download': download(args)
elif args.action == 'inventory': inventory()
else: prepare()
